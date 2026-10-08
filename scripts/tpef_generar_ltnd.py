#!/usr/bin/env python3
"""Generate the TPEF Lemario Total Normalizado y Deduplicado (LTND).

The script reads the 57 glossary worksheets in ``Glosarios_2026.xlsx`` and
exports one record for each normalized lemma, retaining the observed forms and
their source locations. The expected controls are 18,933 source entries,
57 glossaries, 18 countries, and 10,846 LTND forms.
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from collections import OrderedDict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ISO3 = {
    "AR": "ARG", "BO": "BOL", "CL": "CHL", "CO": "COL",
    "CR": "CRI", "CU": "CUB", "DO": "DOM", "EC": "ECU",
    "ES": "ESP", "GT": "GTM", "HN": "HND", "MX": "MEX",
    "NI": "NIC", "PA": "PAN", "PE": "PER", "PY": "PRY",
    "SV": "SLV", "VE": "VEN",
}

APOSTROPHES = str.maketrans({"‘": "'", "’": "'", "‚": "'", "‛": "'"})
QUOTATION_MARKS = str.maketrans({
    "“": '"', "”": '"', "„": '"', "‟": '"', "«": '"', "»": '"'
})


def normalize_lemma(value: object) -> str:
    """Return the historical LTND comparison key."""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = re.sub(r"\s+", " ", text)
    return text.translate(APOSTROPHES).translate(QUOTATION_MARKS)


def ordered_unique(values: list[str]) -> list[str]:
    return list(OrderedDict.fromkeys(values))


def generate_ltnd(input_file: Path) -> pd.DataFrame:
    workbook = load_workbook(input_file, read_only=True, data_only=True)
    groups: dict[str, dict[str, object]] = {}
    source_entries = 0
    countries_seen: set[str] = set()

    for worksheet in workbook.worksheets:
        glossary = worksheet.title
        country = ISO3.get(glossary[:2], glossary[:2])
        countries_seen.add(country)

        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            lemma = row[0] if row else None
            if lemma is None or not str(lemma).strip():
                continue

            display = str(lemma).strip()
            key = normalize_lemma(display)
            typology = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            borrowing = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            source_entries += 1

            group = groups.setdefault(key, {
                "representative": display,
                "forms": [], "glossaries": [], "countries": [],
                "typologies": [], "borrowings": [], "references": [],
                "entry_count": 0,
            })
            group["forms"].append(display)
            group["glossaries"].append(glossary)
            group["countries"].append(country)
            if typology:
                group["typologies"].append(typology)
            if borrowing:
                group["borrowings"].append(borrowing)
            group["references"].append(f"{glossary}!{row_number}")
            group["entry_count"] += 1

    rows = []
    for index, key in enumerate(sorted(groups, key=lambda item: item.casefold()), start=1):
        group = groups[key]
        forms = ordered_unique(group["forms"])
        glossaries = sorted(set(group["glossaries"]))
        countries = sorted(set(group["countries"]))
        typologies = sorted(set(group["typologies"]))
        borrowings = sorted(set(group["borrowings"]))
        rows.append({
            "ID LTND": f"TPEF-LTND-{index:05d}",
            "Lema representativo": group["representative"],
            "Clave normalizada": key,
            "Formas originales observadas": " || ".join(forms),
            "N.º formas originales": len(forms),
            "N.º entradas": group["entry_count"],
            "N.º glosarios": len(glossaries),
            "Glosarios": "; ".join(glossaries),
            "N.º países": len(countries),
            "Países": "; ".join(countries),
            "Tipologías observadas": "; ".join(typologies),
            "Préstamos observados": "; ".join(borrowings),
            "Referencias de origen": "; ".join(group["references"]),
        })

    controls = {
        "entries": source_entries,
        "glossaries": len(workbook.sheetnames),
        "countries": len(countries_seen),
        "ltnd": len(rows),
    }
    expected = {"entries": 18933, "glossaries": 57, "countries": 18, "ltnd": 10846}
    if controls != expected:
        raise ValueError(f"Control totals differ: found {controls}, expected {expected}")
    return pd.DataFrame(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Path to Glosarios_2026.xlsx")
    parser.add_argument("--output-dir", type=Path, default=Path.cwd())
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    dataframe = generate_ltnd(args.input)
    dataframe.to_csv(args.output_dir / "TPEF_LTND_v1.0.csv", index=False, encoding="utf-8-sig")
    dataframe.to_excel(args.output_dir / "TPEF_LTND_v1.0.xlsx", index=False)
    print(f"Generated {len(dataframe):,} LTND forms from 18,933 source entries.")


if __name__ == "__main__":
    main()
