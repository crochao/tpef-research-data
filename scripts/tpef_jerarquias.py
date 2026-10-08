"""
TPEF — Extracción de estructuras conceptuales jerárquicas
=========================================================
Autor: C. O. Rocha Ochoa
Descripción:
    Identifica las estructuras conceptuales jerárquicas latentes en
    el corpus mediante el análisis de la estructura polilexemática de
    los lemas. Dos o más lemas que comparten su primer token forman un
    grupo jerárquico en el que ese token actúa como nodo hiperónimo
    implícito o explícito del conjunto.
    Para cada grupo jerárquico identifica:
      - El nodo padre (hiperónimo): el token compartido
      - Tipo de nodo: EXPLÍCITO si existe como lema independiente
                      en el mismo glosario; IMPLÍCITO en caso contrario
      - Los hipónimos: lemas que comparten el nodo padre
      - Jerarquía de 3 niveles: si algún subconjunto de hipónimos
        comparte además sus dos primeros tokens
    Produce una tabla por glosario y una tabla agregada por país.
Entrada:  Glosarios_2026.xlsx
Salida:   tpef_jerarquias_por_glosario.xlsx
          tpef_jerarquias_resumen.xlsx
          tpef_jerarquias_grupos.xlsx
Requisitos:
    pip install openpyxl pandas
"""
from collections import defaultdict
from openpyxl import load_workbook
import pandas as pd
# ── Configuración ─────────────────────────────────────────────────────────────
INPUT_FILE      = 'Glosarios_2026.xlsx'
OUTPUT_GLOSSARY = 'tpef_jerarquias_por_glosario.xlsx'
OUTPUT_SUMMARY  = 'tpef_jerarquias_resumen.xlsx'
OUTPUT_GROUPS   = 'tpef_jerarquias_grupos.xlsx'
COUNTRY_LABELS = {
    'AR':'Argentina','BO':'Bolivia','CL':'Chile','CO':'Colombia',
    'CR':'Costa Rica','CU':'Cuba','DO':'Rep. Dominicana','EC':'Ecuador',
    'ES':'España','GT':'Guatemala','HN':'Honduras','MX':'México',
    'NI':'Nicaragua','PA':'Panamá','PE':'Perú','PY':'Paraguay',
    'SV':'El Salvador','VE':'Venezuela'
}
# ── Extracción de jerarquías ──────────────────────────────────────────────────
def extract_hierarchies(lemas_list, min_group_size=2):
    """
    Identifica grupos jerárquicos en una lista de lemas.
    Algoritmo:
      1. Para cada lema de más de una palabra, se extrae el primer token
         como candidato a nodo padre.
      2. Se agrupan los lemas que comparten el mismo primer token.
      3. Los grupos con al menos min_group_size miembros constituyen
         una jerarquía conceptual.
      4. Se verifica si el nodo padre existe como lema independiente
         en el glosario (EXPLÍCITO) o no (IMPLÍCITO).
      5. Dentro de cada grupo, se buscan sub-jerarquías de 3 niveles:
         subconjuntos de hipónimos que comparten sus dos primeros tokens.
    Devuelve una lista de grupos jerárquicos con sus propiedades.
    """
    lemas = [l for l in lemas_list if l and l.strip()]
    lemas_lower = {l.lower(): l for l in lemas}
    # Agrupar por primer token
    first_token_groups = defaultdict(list)
    for lema in lemas:
        parts = lema.split()
        if len(parts) >= 2:
            first_token_groups[parts[0].lower()].append(lema)
    groups = []
    for head, children in first_token_groups.items():
        if len(children) < min_group_size:
            continue
        # Tipo de nodo padre
        parent_exists = head in lemas_lower
        parent_lema   = lemas_lower.get(head, None)
        # Buscar sub-jerarquías de 3 niveles
        first_two_groups = defaultdict(list)
        for child in children:
            parts = child.split()
            if len(parts) >= 3:
                key = ' '.join(parts[:2]).lower()
                first_two_groups[key].append(child)
        sub_hierarchies = {
            k: v for k, v in first_two_groups.items()
            if len(v) >= min_group_size
        }
        groups.append({
            'head':            head,
            'parent_lema':     parent_lema,
            'parent_type':     'Explícito' if parent_exists else 'Implícito',
            'children':        sorted(children),
            'n_children':      len(children),
            'has_3_levels':    bool(sub_hierarchies),
            'sub_hierarchies': sub_hierarchies,
        })
    # Ordenar por número de hipónimos descendente
    groups.sort(key=lambda x: -x['n_children'])
    return groups
# ── Script principal ──────────────────────────────────────────────────────────
def main():
    print("TPEF — Extracción de estructuras conceptuales jerárquicas")
    print("="*60)
    wb = load_workbook(INPUT_FILE, read_only=True)
    print(f"Corpus: {len(wb.sheetnames)} glosarios")
    # Preservar orden de países del archivo
    country_order = []
    seen = set()
    for s in wb.sheetnames:
        c = s[:2]
        if c not in seen:
            country_order.append(c)
            seen.add(c)
    rows_glossary = []   # una fila por glosario
    rows_groups   = []   # una fila por grupo jerárquico
    country_agg   = defaultdict(lambda: defaultdict(int))
    corpus_totals = defaultdict(int)
    total_entries = 0
    for sheet_name in wb.sheetnames:
        country = sheet_name[:2]
        ws = wb[sheet_name]
        # Cargar lemas del glosario
        lemas = []
        for row in ws.iter_rows(values_only=True):
            if row[0] and str(row[0]).strip():
                lemas.append(str(row[0]).strip())
                total_entries += 1
        # Extraer jerarquías
        groups = extract_hierarchies(lemas)
        n_groups    = len(groups)
        n_explicit  = sum(1 for g in groups if g['parent_type'] == 'Explícito')
        n_implicit  = sum(1 for g in groups if g['parent_type'] == 'Implícito')
        n_3levels   = sum(1 for g in groups if g['has_3_levels'])
        n_children  = sum(g['n_children'] for g in groups)
        pct_exp     = n_explicit / n_groups * 100 if n_groups else 0
        pct_child   = n_children / len(lemas) * 100 if lemas else 0
        # Fila de resumen por glosario
        rows_glossary.append({
            'Glosario':          sheet_name,
            'País':              COUNTRY_LABELS.get(country, country),
            'Lemas':             len(lemas),
            'Grupos jerárq.':    n_groups,
            'Nodos explícitos':  n_explicit,
            'Nodos implícitos':  n_implicit,
            '% explícitos':      round(pct_exp, 1),
            'Jerarq. 3 niveles': n_3levels,
            'Total hipónimos':   n_children,
            '% lemas en jerarq.': round(pct_child, 1),
        })
        # Filas de grupos detallados
        for g in groups:
            # Sub-jerarquías como texto
            sub_str = '; '.join(
                f"[{k}]: {', '.join(v)}"
                for k, v in g['sub_hierarchies'].items()
            ) if g['sub_hierarchies'] else ''
            rows_groups.append({
                'Glosario':        sheet_name,
                'País':            COUNTRY_LABELS.get(country, country),
                'Nodo padre':      g['head'],
                'Tipo nodo':       g['parent_type'],
                'Lema padre':      g['parent_lema'] or '(implícito)',
                'N hipónimos':     g['n_children'],
                'Hipónimos':       ' | '.join(g['children']),
                '3 niveles':       'Sí' if g['has_3_levels'] else 'No',
                'Sub-jerarquías':  sub_str,
            })
        # Agregar por país
        country_agg[country]['groups']   += n_groups
        country_agg[country]['explicit'] += n_explicit
        country_agg[country]['implicit'] += n_implicit
        country_agg[country]['3levels']  += n_3levels
        country_agg[country]['children'] += n_children
        country_agg[country]['lemas']    += len(lemas)
        # Totales del corpus
        corpus_totals['groups']   += n_groups
        corpus_totals['explicit'] += n_explicit
        corpus_totals['implicit'] += n_implicit
        corpus_totals['3levels']  += n_3levels
        corpus_totals['children'] += n_children
    # ── Resumen por país ──────────────────────────────────────────────────────
    rows_summary = []
    for country in country_order:
        d = country_agg[country]
        tot_g = d['groups']
        rows_summary.append({
            'País':                COUNTRY_LABELS.get(country, country),
            'Lemas':               d['lemas'],
            'Grupos jerárq.':      tot_g,
            'Nodos explícitos':    d['explicit'],
            'Nodos implícitos':    d['implicit'],
            '% explícitos':        round(d['explicit']/tot_g*100, 1) if tot_g else 0,
            'Jerarq. 3 niveles':   d['3levels'],
            'Total hipónimos':     d['children'],
            '% lemas en jerarq.':  round(d['children']/d['lemas']*100, 1)
                                   if d['lemas'] else 0,
        })
    # ── Exportar ──────────────────────────────────────────────────────────────
    pd.DataFrame(rows_glossary).to_excel(OUTPUT_GLOSSARY, index=False)
    pd.DataFrame(rows_summary).to_excel(OUTPUT_SUMMARY, index=False)
    pd.DataFrame(rows_groups).to_excel(OUTPUT_GROUPS, index=False)
    print(f"\nResumen por glosario  → {OUTPUT_GLOSSARY}")
    print(f"Resumen por país      → {OUTPUT_SUMMARY}")
    print(f"Grupos detallados     → {OUTPUT_GROUPS}")
    # ── Estadísticas globales ─────────────────────────────────────────────────
    tot_g  = corpus_totals['groups']
    tot_e  = corpus_totals['explicit']
    tot_i  = corpus_totals['implicit']
    tot_3  = corpus_totals['3levels']
    tot_ch = corpus_totals['children']
    print(f"\nEstadísticas del corpus:")
    print(f"  Total entradas:              {total_entries:,}")
    print(f"  Grupos jerárquicos:          {tot_g:,}")
    print(f"    Nodos explícitos:          {tot_e:,} ({tot_e/tot_g*100:.1f}%)")
    print(f"    Nodos implícitos:          {tot_i:,} ({tot_i/tot_g*100:.1f}%)")
    print(f"    Jerarquías de 3 niveles:   {tot_3:,} ({tot_3/tot_g*100:.1f}%)")
    print(f"  Total hipónimos en grupos:   {tot_ch:,} ({tot_ch/total_entries*100:.1f}% del corpus)")
    # Nodos más productivos del corpus
    print(f"\n  Top 15 nodos más frecuentes (por glosarios en que aparecen):")
    all_heads = defaultdict(lambda: {'glossaries': 0, 'children': 0, 'explicit': 0})
    for row in rows_groups:
        h = row['Nodo padre']
        all_heads[h]['glossaries'] += 1
        all_heads[h]['children']   += row['N hipónimos']
        if row['Tipo nodo'] == 'Explícito':
            all_heads[h]['explicit'] += 1
    for head, d in sorted(
        all_heads.items(), key=lambda x: -x[1]['glossaries']
    )[:15]:
        print(f"    {exp_tag} {head:<25} "
              f"{d['glossaries']:>3} glosarios | "
              f"{d['children']:>4} hipónimos")
    print("\nFinalizado.")
if __name__ == '__main__':
    main()
